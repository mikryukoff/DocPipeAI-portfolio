"""Замена [TABLE_SECTIONS] в паспорте: подставленный «№ п/п» жирный."""
from types import SimpleNamespace

from docx import Document

from backend.app.infrastructure.exporters.passport_builders.sections import fill_sections_table
from backend.app.infrastructure.exporters.passport_builders.placeholders import (
    replace_in_paragraph_bold_token,
    replace_placeholder_in_cell_bold,
)


def test_replace_in_paragraph_bold_token_makes_new_text_bold():
    doc = Document()
    p = doc.add_paragraph("Префикс [X] суффикс")
    assert replace_in_paragraph_bold_token(p, "[X]", "ЗАМЕНА")
    assert p.text == "Префикс ЗАМЕНА суффикс"
    runs = list(p.runs)
    assert runs[0].text == "Префикс "
    assert runs[1].text == "ЗАМЕНА"
    assert runs[1].bold is True
    assert runs[2].text == " суффикс"


def test_replace_placeholder_in_cell_bold():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = "[TABLE_SECTIONS]"
    assert replace_placeholder_in_cell_bold(cell, "[TABLE_SECTIONS]", "№ п/п")
    assert cell.text.strip() == "№ п/п"
    mid = cell.paragraphs[0].runs[0]
    assert mid.text == "№ п/п"
    assert mid.bold is True


def test_fill_sections_table_replaces_placeholder_with_bold():
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "[TABLE_SECTIONS]"
    pipeline = SimpleNamespace(pipelines_len=None)
    assert fill_sections_table(doc, pipeline) is True
    cell = doc.tables[0].rows[0].cells[0]
    assert "№ п/п" in cell.text
    assert cell.paragraphs[0].runs[0].bold is True
