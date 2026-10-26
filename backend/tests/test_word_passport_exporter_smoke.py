import os

import pytest
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn

from backend.app.domain.models.pipeline import (
    FlowDirection,
    Pipeline,
    PipelineLength,
    PipelineParameters,
    VIKEntry,
)
from backend.app.domain.models.specification import Flange, Specification
from backend.app.infrastructure.exporters.passport_builders.spec_tables_builder import (
    SpecTablesBuilder,
)
from backend.app.infrastructure.exporters.word_passport_exporter import WordPassportExporter


def _build_pipeline_for_export() -> Pipeline:
    return Pipeline(
        name="SMOKE-PIPE-001",
        medium="ВОДА",
        mounting_scheme="УЧ-001",
        flow_direction=FlowDirection(start="А", end="Б"),
        parameters=PipelineParameters(
            gost_category="II",
            working_pressure="1.6",
            working_temperature="120",
            calculated_pressure="2.0",
            calculated_temperature="150",
            strength_test_pressure="2.5",
            strength_test_method="ВОДА",
            strength_test_type="гидравлическим",
            strength_test_duration="10 мин",
            strength_test_result="Герметично",
            material="09Г2С",
            installation_date="2026-01-01",
            welding_type="РД",
            welding_material="Св-08Г2С",
        ),
        specifications=[
            Specification(
                page="6.1",
                flanges=[
                    Flange(
                        position=1,
                        name="Фланец smoke",
                        unit="шт",
                        quantity=2,
                        gost="ГОСТ 12820",
                        dn="100",
                        pn="16",
                        steel_grade="09Г2С",
                        flange_gost="ГОСТ 33259",
                        bolt_grade="35Х",
                        bolt_gost="ГОСТ 9066",
                    )
                ],
            )
        ],
        pipelines_len={"DN100": PipelineLength(total_length=10.5, pages=["6.1"])},
    )


def _require_word_template() -> str:
    template_path = os.getenv("WORD_TEMPLATE_PATH")
    if not template_path or not os.path.exists(template_path):
        pytest.skip("Word template omitted in portfolio snapshot (NDA)")
    return template_path


def test_exported_table_row_has_cant_split_and_centering_policy(tmp_path):
    template_path = _require_word_template()

    exporter = WordPassportExporter(template_path=template_path)
    pipeline = _build_pipeline_for_export()
    output_path = tmp_path / "smoke_passport.docx"

    exporter.export(pipeline, output_path)
    assert output_path.exists()

    doc = Document(output_path)
    target_row = None
    for table in doc.tables:
        for row in table.rows:
            if any("Фланец smoke" in cell.text for cell in row.cells):
                target_row = row
                break
        if target_row is not None:
            break

    assert target_row is not None, "В сгенерированном паспорте не найдена строка с тестовым фланцем"
    tr_pr = target_row._tr.get_or_add_trPr()
    assert tr_pr.find(qn("w:cantSplit")) is not None
    assert target_row.height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST


def test_strength_test_mapping_hydraulic_only():
    exporter = WordPassportExporter.__new__(WordPassportExporter)
    pipeline = _build_pipeline_for_export()
    pipeline.parameters.strength_test_type = "гидравлическим"

    mapping = exporter._build_strength_test_mapping(pipeline)

    assert "ДАВЛЕНИЕ ГИДРАВЛИЧЕСКОГО ИСПЫТАНИЯ" in mapping
    assert mapping["ДАВЛЕНИЕ ПНЕВМАТИЧЕСКОГО ИСПЫТАНИЯ"] == "---"
    assert mapping["СРЕДА ПНЕВМАТИЧЕСКОГО ИСПЫТАНИЯ"] == "---"
    assert mapping["ПРОДОЛЖИТЕЛЬНОСТЬ ПНЕВМАТИЧЕСКОЙ ВЫДЕРЖКИ"] == "---"


def test_strength_test_mapping_pneumatic_only():
    exporter = WordPassportExporter.__new__(WordPassportExporter)
    pipeline = _build_pipeline_for_export()
    pipeline.parameters.strength_test_type = "пневматическим"

    mapping = exporter._build_strength_test_mapping(pipeline)

    assert "ДАВЛЕНИЕ ПНЕВМАТИЧЕСКОГО ИСПЫТАНИЯ" in mapping
    assert mapping["ДАВЛЕНИЕ ГИДРАВЛИЧЕСКОГО ИСПЫТАНИЯ"] == "---"
    assert mapping["СРЕДА ГИДРАВЛИЧЕСКОГО ИСПЫТАНИЯ"] == "---"
    assert mapping["ПРОДОЛЖИТЕЛЬНОСТЬ ГИДРАВЛИЧЕСКОЙ ВЫДЕРЖКИ"] == "---"


def test_strength_test_mapping_both_types_fill_both():
    exporter = WordPassportExporter.__new__(WordPassportExporter)
    pipeline = _build_pipeline_for_export()
    pipeline.parameters.strength_test_type = "гидравлическим и пневматическим"

    mapping = exporter._build_strength_test_mapping(pipeline)

    assert "ДАВЛЕНИЕ ГИДРАВЛИЧЕСКОГО ИСПЫТАНИЯ" in mapping
    assert "ДАВЛЕНИЕ ПНЕВМАТИЧЕСКОГО ИСПЫТАНИЯ" in mapping


def test_strength_test_pressure_format_from_dot_number():
    exporter = WordPassportExporter.__new__(WordPassportExporter)
    pipeline = _build_pipeline_for_export()
    pipeline.parameters.strength_test_type = "гидравлическим"
    pipeline.parameters.strength_test_pressure = "1.17 МПа (11.70 кгс/см²)"

    mapping = exporter._build_strength_test_mapping(pipeline)

    assert mapping["ДАВЛЕНИЕ ГИДРАВЛИЧЕСКОГО ИСПЫТАНИЯ"] == "1,17 МПа (11,70 кгс/см²)"


def test_strength_test_pressure_format_autofills_kgf_when_missing():
    exporter = WordPassportExporter.__new__(WordPassportExporter)
    pipeline = _build_pipeline_for_export()
    pipeline.parameters.strength_test_type = "гидравлическим"
    pipeline.parameters.strength_test_pressure = "1.7 МПа (кгс/см²)"

    mapping = exporter._build_strength_test_mapping(pipeline)

    assert mapping["ДАВЛЕНИЕ ГИДРАВЛИЧЕСКОГО ИСПЫТАНИЯ"] == "1,70 МПа (17,00 кгс/см²)"


def test_spec_tables_builder_build_pipe_scheme_text_deduplicates_number_sign():
    builder = SpecTablesBuilder()

    assert builder._build_pipe_scheme_text(
        medium="ВОДА",
        medium_genitive="воды",
        pipe_num="№ЗВ.1",
    ) == "Схема сварных стыков трубопровода воды №ЗВ.1"

    assert builder._build_pipe_scheme_text(
        medium="",
        medium_genitive=None,
        pipe_num="№№41001",
    ) == "Схема сварных стыков трубопровода №41001"


def test_spec_tables_builder_filters_vik_rows_by_type():
    rows = [
        VIKEntry(pipe_num="41001", statement_text="001-01.01.26", type="ВИК"),
        VIKEntry(pipe_num="41002", statement_text="002-01.01.26", type="УЗК"),
        VIKEntry(pipe_num="41003", statement_text="003-01.01.26", type="ПВК"),
    ]

    builder = SpecTablesBuilder()

    assert len(builder._filter_vik_rows_by_type(rows, "ВИК")) == 1
    assert len(builder._filter_vik_rows_by_type(rows, "УЗК")) == 1
    assert len(builder._filter_vik_rows_by_type(rows, "ПВК")) == 1


def test_spec_tables_builder_fill_table_45_combines_vik_uzk_pvk():
    doc = Document()
    table = doc.add_table(rows=2, cols=7)
    for idx in range(7):
        table.rows[1].cells[idx].text = f"T{idx + 1}"

    rows = [
        VIKEntry(pipe_num="41001", statement_text="001-01.01.26", type="ВИК"),
        VIKEntry(pipe_num="41002", statement_text="002-01.01.26", type="УЗК"),
    ]

    builder = SpecTablesBuilder()
    builder._fill_table_4_5_combined(
        table,
        rows,
        medium="ВОДА",
        gost_category="Б II",
    )

    assert len(table.rows) == 4

    vik_row = table.rows[1]
    assert vik_row.cells[0].text == "T1"
    assert vik_row.cells[3].text == "T4"
    assert "№41001" in vik_row.cells[1].text
    assert vik_row.cells[2].text == "№001-01.01.26"
    assert vik_row.cells[4].text == "100%"
    assert vik_row.cells[5].text == "Не обнаружено"
    assert vik_row.cells[6].text == "Годен"

    uzk_row = table.rows[2]
    assert uzk_row.cells[0].text == "T1"
    assert uzk_row.cells[3].text == "T4"
    assert uzk_row.cells[2].text == "002-01.01.26"
    assert uzk_row.cells[4].text == "20%"
    assert uzk_row.cells[5].text == "Не обнаружено"
    assert uzk_row.cells[6].text == "Годен"

    pvk_row = table.rows[3]
    assert pvk_row.cells[0].text == "T1"
    assert pvk_row.cells[3].text == "T4"
    assert pvk_row.cells[1].text == "---"
    assert pvk_row.cells[2].text == "---"
    assert pvk_row.cells[4].text == "---"
    assert pvk_row.cells[5].text == "---"
    assert pvk_row.cells[6].text == "---"


def test_spec_tables_builder_fill_table_45_keeps_first_fourth_when_vik_missing():
    doc = Document()
    table = doc.add_table(rows=2, cols=7)
    for idx in range(7):
        table.rows[1].cells[idx].text = f"M{idx + 1}"

    builder = SpecTablesBuilder()
    builder._fill_table_4_5_combined(
        table,
        vik_rows=[],
        medium="ВОДА",
        gost_category="Аа(1)",
    )

    vik_row = table.rows[1]
    assert vik_row.cells[0].text == "M1"
    assert vik_row.cells[3].text == "M4"
    assert vik_row.cells[1].text == "---"
    assert vik_row.cells[2].text == "---"
    assert vik_row.cells[4].text == "---"
    assert vik_row.cells[5].text == "---"
    assert vik_row.cells[6].text == "---"


def test_spec_tables_builder_fill_table_46_styloscopy():
    doc = Document()
    table = doc.add_table(rows=2, cols=6)
    for idx in range(6):
        table.rows[1].cells[idx].text = f"S{idx + 1}"

    rows = [
        VIKEntry(
            pipe_num="YY03-90454-CSC-25-E010H2A00-FC",
            statement_text='{"protocol_after_no":"ОК 1510/4-25 от 15 октября 2025г.","steel_marks":"EN 1.4404, EN 1.4571"}',
            type="СТИЛОСКОПИРОВАНИЕ",
        )
    ]
    builder = SpecTablesBuilder()
    builder._fill_table_4_6_styloscopy(
        table,
        rows,
        medium="ВОДА",
        gost_category="Б II",
    )

    row = table.rows[1]
    assert row.cells[0].text == "S1"
    assert "YY03-90454-CSC-25-E010H2A00-FC" in row.cells[1].text
    assert row.cells[2].text == "№ОК 1510/4-25 от 15 октября 2025г."
    assert row.cells[3].text == "S4"
    assert row.cells[4].text == "100%"
    assert row.cells[5].text == "EN 1.4404, EN 1.4571"
