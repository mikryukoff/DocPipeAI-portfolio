from types import SimpleNamespace

from docx import Document

from backend.app.infrastructure.exporters.passport_builders.spec_tables_builder import (
    SpecTablesBuilder,
)
from backend.app.infrastructure.exporters.passport_builders.xml_helpers import (
    apply_keep_with_next_before_table,
    copy_table_after,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def test_fill_table_4_3_keeps_two_header_rows_and_writes_data_after_them():
    doc = Document()
    table = doc.add_table(rows=2, cols=10)

    # Двухстрочный заголовок должен сохраниться без изменений.
    table.rows[0].cells[6].text = "Материал Фланцев"
    table.rows[0].cells[8].text = "Материал шпилек, гаек и болтов"
    table.rows[1].cells[6].text = "марка стали"
    table.rows[1].cells[7].text = "ГОСТ,ТУ"
    table.rows[1].cells[8].text = "марка стали"
    table.rows[1].cells[9].text = "ГОСТ,ТУ"

    spec = SimpleNamespace(
        flanges=[
            SimpleNamespace(
                name="Фланец 1",
                quantity=2,
                gost="ГОСТ 123",
                dn="100",
                pn="16",
                steel_grade="09Г2С",
                flange_gost="ГОСТ 33259",
                bolt_grade="35Х",
                bolt_gost="ГОСТ 9066",
            )
        ],
        supports=[],
    )

    builder = SpecTablesBuilder()

    filled = builder._fill_table_4_3(table, spec)

    assert filled is True
    assert len(table.rows) == 3

    # Проверяем, что строки заголовка не затёрлись.
    assert table.rows[0].cells[6].text == "Материал Фланцев"
    assert table.rows[0].cells[8].text == "Материал шпилек, гаек и болтов"
    assert table.rows[1].cells[6].text == "марка стали"
    assert table.rows[1].cells[7].text == "ГОСТ,ТУ"
    assert table.rows[1].cells[8].text == "марка стали"
    assert table.rows[1].cells[9].text == "ГОСТ,ТУ"

    # Данные должны начинаться только после двух строк заголовка.
    data_row = table.rows[2]
    assert data_row.cells[0].text == "1"
    assert data_row.cells[1].text == "Фланец 1"
    assert data_row.cells[2].text == "2"
    assert data_row.cells[3].text == "ГОСТ 123"
    assert data_row.cells[4].text == "100"
    assert data_row.cells[5].text == "16"
    assert data_row.cells[6].text == "09Г2С"
    assert data_row.cells[7].text == "ГОСТ 33259"
    assert data_row.cells[8].text == "35Х"
    assert data_row.cells[9].text == "ГОСТ 9066"


def test_fill_table_4_3_replaces_empty_values_with_dashes():
    doc = Document()
    table = doc.add_table(rows=2, cols=10)

    spec = SimpleNamespace(
        flanges=[
            SimpleNamespace(
                name="Фланец 2",
                quantity=1,
                gost="",
                dn="100",
                pn="16",
                steel_grade="",
                flange_gost="",
                bolt_grade="",
                bolt_gost="",
            )
        ],
        supports=[],
    )

    builder = SpecTablesBuilder()
    filled = builder._fill_table_4_3(table, spec)

    assert filled is True
    data_row = table.rows[2]
    assert data_row.cells[3].text == "---"
    assert data_row.cells[6].text == "---"
    assert data_row.cells[7].text == "---"
    assert data_row.cells[8].text == "---"
    assert data_row.cells[9].text == "---"


def test_copy_table_after_preserves_requested_header_rows():
    doc = Document()
    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "H1"
    table.rows[1].cells[0].text = "H2"
    table.rows[2].cells[0].text = "D1"
    table.rows[3].cells[0].text = "D2"

    copied = copy_table_after(table, "6.2", header_rows=2)

    assert len(copied.rows) == 2
    assert copied.rows[0].cells[0].text == "H1"
    assert copied.rows[1].cells[0].text == "H2"


def test_apply_keep_with_next_before_table_marks_previous_paragraphs():
    doc = Document()
    p1 = doc.add_paragraph("4.1 Сведения о трубах")
    p2 = doc.add_paragraph("")
    p3 = doc.add_paragraph("Участок 6 (Схема сварных соединений #6)")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "x"

    # Имитируем случайный признак разрыва страницы в абзаце шаблона.
    p1_pr = p1._p.get_or_add_pPr()
    p1_pr.append(OxmlElement("w:pageBreakBefore"))

    apply_keep_with_next_before_table(table, max_paragraphs=20)

    p1_pr = p1._p.find(qn("w:pPr"))
    p2_pr = p2._p.find(qn("w:pPr"))
    p3_pr = p3._p.find(qn("w:pPr"))
    assert p1_pr is not None and p1_pr.find(qn("w:keepNext")) is not None
    assert p2_pr is not None and p2_pr.find(qn("w:keepNext")) is not None
    assert p3_pr is not None and p3_pr.find(qn("w:keepNext")) is not None
    assert p1_pr.find(qn("w:pageBreakBefore")) is None


def test_sort_pages_handles_plain_numbers():
    builder = SpecTablesBuilder()
    pages = ["10", "2", "1", "A"]

    assert builder._sort_pages(pages) == ["1", "2", "10", "A"]


def test_fill_table_4_4_embedments_autofills_location_from_pipeline_name():
    doc = Document()
    table = doc.add_table(rows=1, cols=7)

    spec = SimpleNamespace(
        instrument_embedments=[
            SimpleNamespace(
                name="ЗК-1",
                location="---",
                dn="50",
                pn="16",
                material_grade="09Г2С",
                gost="ГОСТ 1",
                embedment_elements=[
                    {
                        "name": "Элемент 1",
                        "location": "---",
                        "dn": "50",
                        "pn": "16",
                        "material_grade": "09Г2С",
                        "gost": "ГОСТ 1",
                    },
                    {
                        "name": "Элемент 2",
                        "location": "Блок Б",
                        "dn": "80",
                        "pn": "25",
                        "material_grade": "12Х18Н10Т",
                        "gost": "ГОСТ 2",
                    },
                ],
            )
        ]
    )

    builder = SpecTablesBuilder()
    filled = builder._fill_table_4_4_embedments(
        table,
        spec,
        pipeline_name="41079",
    )

    assert filled is True
    assert table.rows[1].cells[2].text == "т-д 41079"
    assert table.rows[2].cells[2].text == "т-д 41079"


def test_placeholder_replacement_uses_bold_run_not_cell_text():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = "[4_1]"

    SpecTablesBuilder._replace_placeholder_in_cell_with_bold(cell, "[4_1]", "№ п/п")

    assert cell.text.strip() == "№ п/п"
    para = cell.paragraphs[0]
    assert any(r.text == "№ п/п" and r.bold for r in para.runs)


def test_placeholder_split_across_runs_still_replaced_bold():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    p.add_run("Префикс [4_")
    p.add_run("1] суффикс")

    SpecTablesBuilder._replace_placeholder_in_cell_with_bold(cell, "[4_1]", "№ п/п")

    assert "№ п/п" in cell.text
    assert "Префикс " in cell.text
    assert " суффикс" in cell.text
    assert any(r.text == "№ п/п" and r.bold for r in p.runs)
