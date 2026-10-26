"""Тесты извлечения ВИК и испытаний из DOCX без БД."""

from docx import Document
from docx.enum.text import WD_BREAK

from backend.app.infrastructure.files.docx.docx_logical_pages import iter_logical_pages
from backend.app.infrastructure.files.docx.pipelines_tests_docx_extractor import (
    extract_pipelines_tests_from_logical_page,
    extract_pipelines_tests_rows_from_document,
)
from backend.app.infrastructure.files.docx.vik_docx_extractor import (
    extract_vik_from_logical_page,
    extract_vik_rows_from_document,
)


def _add_page_break_after(paragraph):
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def test_vik_single_page_paragraph_and_table():
    doc = Document()
    doc.add_paragraph("ЗАКЛЮЧЕНИЕ 12-34.56.78 дополнительно")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Конструкция, узел:"
    t.rows[0].cells[1].text = "AB12-34567.89-CDE-12-FGH"

    pages = iter_logical_pages(doc)
    assert len(pages) == 1
    row = extract_vik_from_logical_page(pages[0])
    assert row is not None
    assert row["statement_text"] == "12-34.56.78"
    assert row["pipe_num"] == "AB12-34567.89-CDE-12-FGH"


def test_vik_merged_like_row_skips_duplicate_header_cell():
    doc = Document()
    doc.add_paragraph("ЗАКЛЮЧЕНИЕ 22-07.08.25")
    t = doc.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text = "Конструкция, узел:"
    # Имитируем поведение merged-ячейки: дублирование заголовка в соседней колонке.
    t.rows[0].cells[1].text = "Конструкция, узел:"
    t.rows[0].cells[2].text = "95675"

    row = extract_vik_rows_from_document(doc)[0]
    assert row["statement_text"] == "22-07.08.25"
    assert row["pipe_num"] == "95675"


def test_vik_two_logical_pages():
    doc = Document()
    doc.add_paragraph("ЗАКЛЮЧЕНИЕ 11-22.33.44 первая дата")
    t1 = doc.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text = "Конструкция, узел:"
    t1.rows[0].cells[1].text = "LINE-A"
    p = doc.add_paragraph()
    _add_page_break_after(p)
    doc.add_paragraph("ЗАКЛЮЧЕНИЕ 55-66.77.88 вторая дата")
    t2 = doc.add_table(rows=1, cols=2)
    t2.rows[0].cells[0].text = "Конструкция, узел:"
    t2.rows[0].cells[1].text = "LINE-B"

    rows = extract_vik_rows_from_document(doc)
    assert len(rows) == 2
    assert rows[0]["pipe_num"] == "LINE-A"
    assert rows[0]["statement_text"] == "11-22.33.44"
    assert rows[1]["pipe_num"] == "LINE-B"
    assert rows[1]["statement_text"] == "55-66.77.88"


def test_pipelines_tests_paragraphs():
    doc = Document()
    doc.add_paragraph("Всего участков трубопровода 5")
    doc.add_paragraph("линия XX00-00000.00-ABC-01-NAME")
    doc.add_paragraph("Продолжительность испытания 30 мин")

    pages = iter_logical_pages(doc)
    assert len(pages) == 1
    row = extract_pipelines_tests_from_logical_page(pages[0])
    assert row is not None
    assert row["pipeline"] == "XX00-00000.00-ABC-01-NAME"
    assert row["test_time"] == "30"


def test_pipelines_tests_duration_on_next_line():
    doc = Document()
    doc.add_paragraph("участков трубопровода")
    doc.add_paragraph("линия PIPE-1")
    doc.add_paragraph("Продолжительность испытания")
    doc.add_paragraph("45")

    row = extract_pipelines_tests_rows_from_document(doc)[0]
    assert row["test_time"] == "45"
    assert row["pipeline"] == "PIPE-1"


def test_pipelines_tests_table_rows():
    doc = Document()
    t = doc.add_table(rows=3, cols=1)
    t.rows[0].cells[0].text = "участков трубопровода 3"
    t.rows[1].cells[0].text = "линия Z-9"
    t.rows[2].cells[0].text = "Продолжительность испытания 60"

    row = extract_pipelines_tests_rows_from_document(doc)[0]
    assert row["pipeline"] == "Z-9"
    assert row["test_time"] == "60"
